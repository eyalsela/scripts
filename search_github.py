import concurrent.futures
import os

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

GITHUB_TOKEN = os.getenv("GITHUB_API_TOKEN")

if not GITHUB_TOKEN:
    raise EnvironmentError("GitHub token not found. Please set the GITHUB_API_TOKEN environment variable.")

headers = {"Authorization": f"token {GITHUB_TOKEN}"}

search_query = "(LLM OR LLMS OR chatgpt OR openai) in:readme,title,description"

repos = []


def fetch_page(page):
    params = {"q": search_query, "sort": "stars", "order": "desc", "per_page": 100, "page": page}
    response = requests.get("https://api.github.com/search/repositories", params=params, headers=headers)
    if response.status_code != 200:
        print(f"Error fetching page {page}: {response.status_code}")
        print(response.text)
        return []
    data = response.json()
    items = data.get("items", [])
    print(f"Fetched page {page}, {len(items)} repositories.")
    return items


# Initial request to get the total count
params = {"q": search_query, "sort": "stars", "order": "desc", "per_page": 100, "page": 1}
response = requests.get("https://api.github.com/search/repositories", params=params, headers=headers)
if response.status_code != 200:
    print(f"Error fetching page 1: {response.status_code}")
    print(response.text)
    total_count = 0
else:
    data = response.json()
    total_count = min(data.get("total_count", 0), 1000)  # GitHub API limits to 1000 results
    print(f"Total repositories found: {total_count}")

    repos.extend(data.get("items", []))
    print(f'Fetched page 1, {len(data.get("items", []))} repositories.')

    # Calculate the number of pages needed
    pages = (total_count - 1) // 100 + 1

    # Fetch the rest of the pages using threads
    with concurrent.futures.ThreadPoolExecutor() as executor:
        futures = [executor.submit(fetch_page, page) for page in range(2, pages + 1)]
        for future in concurrent.futures.as_completed(futures):
            repos.extend(future.result())

print(f"Total repositories fetched: {len(repos)}")

# Prepare the data for the DataFrame
repo_list = []
for repo in repos:
    repo_data = {
        "id": repo["id"],
        "name": repo["name"],
        "full_name": repo["full_name"],
        "description": repo["description"],
        "stargazers_count": repo.get("stargazers_count", "N/A"),
        "forks_count": repo.get("forks_count", "N/A"),
        "language": repo.get("language", "N/A"),
        "updated_at": repo.get("updated_at", "N/A"),
        "owner_login": repo["owner"]["login"],
        "owner_html_url": repo["owner"]["html_url"],
        "created_at": repo["created_at"],
        "pushed_at": repo["pushed_at"],
        "html_url": repo["html_url"],
    }
    repo_list.append(repo_data)

df = pd.DataFrame(repo_list)

# Save to Excel file
while True:
    try:
        df.to_excel("github_repos.xlsx", index=False)
        print("Data saved to github_repos.xlsx")

        # Load the workbook and select the active worksheet
        wb = load_workbook("github_repos.xlsx")
        ws = wb.active

        # Define the table
        tab = Table(displayName="RepoTable", ref=ws.dimensions)

        # Add a default style with striped rows and banded columns
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=True,
        )
        tab.tableStyleInfo = style

        # Add the table to the worksheet
        ws.add_table(tab)

        # Save the workbook
        wb.save("github_repos.xlsx")
        break
    except PermissionError:
        input("Permission denied: Please close 'github_repos.xlsx' and press Enter to try again.")
